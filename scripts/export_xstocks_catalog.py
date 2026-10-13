#!/usr/bin/env python3
"""Write a static xStocks catalog workbook from data/xstocks_book.sqlite.

Prices are not included. Multipliers come from the issuer's public endpoint,
one network per symbol, and are cached in data/xstocks_multipliers.json.
"""
from __future__ import annotations

import json
import sqlite3
import sys
import time
import urllib.parse
import urllib.request
from concurrent.futures import ThreadPoolExecutor, as_completed
from decimal import Decimal
from pathlib import Path
from urllib.error import HTTPError, URLError

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
DB = ROOT / "data" / "xstocks_book.sqlite"
CACHE = ROOT / "data" / "xstocks_multipliers.json"
OUT = ROOT / "data" / "xstocks_catalog_2026-09-26.xlsx"
PROXY = "http://127.0.0.1:4780"

# Issuer API network names are not the same strings as our canonical names.
API_NETWORK = {
    "Ethereum": "Ethereum",
    "Solana": "Solana",
    "XLayer": "XLayer",
    "Arbitrum": "Arbitrum",
    "BNB Smart Chain": "BinanceSmartChain",
    "Optimism": "Optimism",
    "Mantle": "Mantle",
    "Ink": "Ink",
    "HyperEVM": "HyperEVM",
    "TON": "TON",
    "Tron": "Tron",
}
PREFER = [
    "Ethereum", "Solana", "XLayer", "Arbitrum", "BNB Smart Chain",
    "Optimism", "Mantle", "Ink", "HyperEVM", "TON", "Tron",
]
CHAIN_ORDER = PREFER

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
ONLINE_FILL = PatternFill("solid", fgColor="C6EFCE")
POST_FILL = PatternFill("solid", fgColor="FFE2B8")
NONE_FILL = PatternFill("solid", fgColor="F2F2F2")
WRAP_FILL = PatternFill("solid", fgColor="DDEBF7")
TITLE_FONT = Font(bold=True, size=14)
WRAP = Alignment(wrap_text=True, vertical="center")


def load_rows():
    conn = sqlite3.connect(f"file:{DB}?mode=ro", uri=True)
    conn.row_factory = sqlite3.Row
    under = list(conn.execute(
        "SELECT * FROM underlyings WHERE missing_since_ms IS NULL ORDER BY issuer_symbol"))
    deps = list(conn.execute(
        """SELECT * FROM deployments WHERE missing_since_ms IS NULL
           ORDER BY issuer_symbol, network, form"""))
    exits = list(conn.execute(
        "SELECT * FROM spot_exits WHERE missing_since_ms IS NULL ORDER BY pair"))
    conn.close()
    return under, deps, exits


def fetch_multiplier(symbol: str, api_network: str) -> dict:
    url = (
        "https://api.xstocks.fi/api/v2/public/assets/"
        f"{urllib.parse.quote(symbol)}/multiplier?network={urllib.parse.quote(api_network)}"
    )
    request = urllib.request.Request(url, headers={"User-Agent": "xstocks-book/0.1"})
    opener = urllib.request.build_opener(
        urllib.request.ProxyHandler({"http": PROXY, "https": PROXY}))
    with opener.open(request, timeout=20) as response:
        payload = json.loads(response.read().decode(), parse_float=Decimal)
    current = payload.get("currentMultiplier")
    pending = payload.get("newMultiplier")
    return {
        "network": api_network,
        "current": None if current in (None, "") else format(Decimal(str(current)), "f"),
        "pending": None if pending in (None, "", 0, "0") else format(Decimal(str(pending)), "f"),
        "activation": payload.get("activationDateTime") or None,
        "reason": payload.get("reason"),
        "error": None,
    }


def multipliers_for(symbols_networks: dict[str, list[str]], cache: dict) -> dict:
    pending = []
    for symbol, networks in symbols_networks.items():
        cached = cache.get(symbol)
        if cached and cached.get("current") and not cached.get("error"):
            continue
        api_names = [API_NETWORK[name] for name in networks if name in API_NETWORK]
        if not api_names:
            cache[symbol] = {"network": None, "current": None, "pending": None,
                             "activation": None, "reason": None,
                             "error": "no mapped network"}
            continue
        pending.append((symbol, api_names))

    def one(item):
        symbol, api_names = item
        last = "no attempt"
        for api_name in api_names[:2]:
            try:
                row = fetch_multiplier(symbol, api_name)
                return symbol, row
            except (HTTPError, URLError, TimeoutError, OSError, json.JSONDecodeError) as exc:
                last = str(exc)
        return symbol, {"network": api_names[0], "current": None, "pending": None,
                        "activation": None, "reason": None, "error": last}

    done = 0
    with ThreadPoolExecutor(max_workers=4) as pool:
        futures = [pool.submit(one, item) for item in pending]
        for future in as_completed(futures):
            symbol, row = future.result()
            cache[symbol] = row
            done += 1
            if done % 100 == 0 or done == len(pending):
                CACHE.write_text(json.dumps(cache, ensure_ascii=False))
                print(f"multipliers {done}/{len(pending)}", flush=True)
    CACHE.write_text(json.dumps(cache, ensure_ascii=False))
    return cache


def style_header(ws, ncol):
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ncol)}{ws.max_row}"
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 32
    ws.auto_filter.ref = f"A1:{get_column_letter(ncol)}{max(ws.max_row, 2)}"
    ws.sheet_view.showGridLines = False


def widths(ws, sizes):
    for idx, size in enumerate(sizes, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = size


def status_fill(status: str | None):
    if status == "online":
        return ONLINE_FILL
    if status == "post_only":
        return POST_FILL
    return NONE_FILL


def sell_text(status: str | None) -> str:
    if status == "online":
        return "可以。市价或吃买盘都能卖"
    if status == "post_only":
        return "不能直接吃买盘。只能挂不立刻成交的限价卖单"
    if status:
        return status
    return "Kraken 没有这个现货"


def main() -> int:
    under, deps, exits = load_rows()
    by_symbol_deps: dict[str, list] = {}
    for row in deps:
        by_symbol_deps.setdefault(row["issuer_symbol"], []).append(row)
    by_symbol_exit = {row["issuer_symbol"]: row for row in exits if row["issuer_symbol"]}
    networks_for = {
        symbol: [name for name in PREFER if any(d["network"] == name for d in rows)]
        + sorted({d["network"] for d in rows} - set(PREFER))
        for symbol, rows in by_symbol_deps.items()
    }
    cache = json.loads(CACHE.read_text()) if CACHE.exists() else {}
    print(f"fetching multipliers, cached {sum(1 for v in cache.values() if v.get('current'))}", flush=True)
    cache = multipliers_for(networks_for, cache)

    chains = [name for name in CHAIN_ORDER if any(d["network"] == name for d in deps)]
    extra_chains = sorted({d["network"] for d in deps} - set(chains))
    chains.extend(extra_chains)

    wb = Workbook()

    readme = wb.active
    readme.title = "阅读说明"
    notes = [
        ("这张表是什么", "xStocks 发行目录和 Kraken 现货状态。不含价格。来自 data/xstocks_book.sqlite，2026-09-26 同步。"),
        ("库有没有在跑", "没有。采样进程不在运行。文件大约 23MB。目录本身不会把硬盘写满。只有把 15 分钟采样开起来，观察表才会按天增长。"),
        ("股票有多少", f"发行方在册 {len(under)} 只。Kraken 现货对得上的 {len(by_symbol_exit)} 只。"),
        ("online 和 post_only", "这是 Kraken 交易对自己的 status，不是正股开市/休市。online：可以市价、也可以吃掉对手盘。post_only：市场只接受挂单（maker），不能直接打到买一上卖掉。依据是 Kraken 错误说明 EService:Market in post_only mode（docs.kraken.com/exchange/guides/general/errors）。12 与 165 是 2026-09-26 接口实数，不是某份文档里的名单。当天是周六，AAPL 等 12 个仍是 online，所以不是「休市就不能卖」。"),
        ("multiplier", "资产列表里没有这个字段。官方说明在 docs.xstocks.fi/developers/multipliers：起始 1 枚代币 = 1 股，分红、拆股、合股之后用 multiplier 调整。接口是 GET /api/v2/public/assets/{符号}/multiplier?network=。本表每个符号查了一条它实际部署的网络。以太坊和 Solana 上的 AAPLx 实测是同一个数。"),
        ("wrapper", "「有包装代币」指发行方登记了 wrapperAddressV2。原生和包装是两个合约。表里没有单独的「web 端」字段；Kraken 上 177 个交易对网站和 API 都有条目。"),
        ("怎么筛", "股票清单、Kraken现货、部署明细 都开了筛选。Kraken状态 一列：绿色 online，橙色 post_only。"),
    ]
    readme["A1"] = "xStocks 目录说明"
    readme["A1"].font = TITLE_FONT
    readme.append(["项目", "说明"])
    for cell in readme[2]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for title, body in notes:
        readme.append([title, body])
    widths(readme, [24, 120])
    for row in readme.iter_rows(min_row=3, max_col=2):
        row[1].alignment = WRAP
        readme.row_dimensions[row[0].row].height = 48
    readme.freeze_panes = "A3"
    readme.sheet_view.showGridLines = False

    countries = {}
    for row in under:
        countries[row["listing_country"] or "未知"] = countries.get(row["listing_country"] or "未知", 0) + 1
    online = sum(1 for row in exits if row["status"] == "online")
    post_only = sum(1 for row in exits if row["status"] == "post_only")
    with_wrap = 0
    for rows in by_symbol_deps.values():
        if any(d["form"] == "wrapper_v2" for d in rows):
            with_wrap += 1
    chain_counts = []
    for chain in chains:
        native = sum(1 for d in deps if d["network"] == chain and d["form"] == "native")
        wrap = sum(1 for d in deps if d["network"] == chain and d["form"] == "wrapper_v2")
        chain_counts.append((chain, native, wrap))

    summary = wb.create_sheet("总览")
    summary["A1"] = "总览"
    summary["A1"].font = TITLE_FONT
    summary.append([])
    summary.append(["项目", "数量"])
    for cell in summary[3]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for label, value in [
        ("发行方股票/ETF", len(under)),
        ("其中美国", countries.get("US", 0)),
        ("英国", countries.get("GB", 0)),
        ("香港", countries.get("HK", 0)),
        ("德国", countries.get("DE", 0)),
        ("西班牙", countries.get("ES", 0)),
        ("Kraken 有现货", len(by_symbol_exit)),
        ("Kraken online，可吃单卖出", online),
        ("Kraken post_only，只能挂单", post_only),
        ("发行方有、Kraken 没有现货", len(under) - len(by_symbol_exit)),
        ("至少一条链有 wrapper", with_wrap),
        ("部署行（链 × 形态）", len(deps)),
    ]:
        summary.append([label, value])
    summary.append([])
    summary.append(["链", "原生代币数", "wrapper 数"])
    for cell in summary[summary.max_row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for chain, native, wrap in chain_counts:
        summary.append([chain, native, wrap])
    widths(summary, [36, 18, 16])
    summary.freeze_panes = "A4"
    summary.sheet_view.showGridLines = False

    stock = wb.create_sheet("股票清单")
    stock_headers = [
        "发行方符号", "正股代码", "上市地", "代币名称", "代币ISIN", "正股ISIN",
        "Kraken交易对", "Kraken状态", "能不能直接卖给买盘",
        "部署链数", "有wrapper", "wrapper所在链", "仅原生的链",
        "multiplier查询网络", "当前multiplier", "待生效multiplier", "multiplier备注",
        "最小申赎美元",
    ] + [f"{chain}" for chain in chains]
    stock.append(stock_headers)
    def rank(row):
        exit_row = by_symbol_exit.get(row["issuer_symbol"])
        if exit_row is not None and exit_row["status"] == "online":
            group = 0
        elif exit_row is not None:
            group = 1
        else:
            group = 2
        return group, row["issuer_symbol"]

    ranked = sorted(under, key=rank)
    for row in ranked:
        symbol = row["issuer_symbol"]
        dep_rows = by_symbol_deps.get(symbol, [])
        exit_row = by_symbol_exit.get(symbol)
        per_chain = {}
        for dep in dep_rows:
            per_chain.setdefault(dep["network"], set()).add(dep["form"])
        wrap_chains = [chain for chain in chains if "wrapper_v2" in per_chain.get(chain, ())]
        native_only = [chain for chain in chains
                       if per_chain.get(chain) == {"native"}]
        mult = cache.get(symbol) or {}
        values = [
            symbol,
            row["underlying_symbol"],
            row["listing_country"],
            row["name"],
            row["token_isin"],
            row["underlying_isin"],
            None if exit_row is None else exit_row["pair"],
            None if exit_row is None else exit_row["status"],
            sell_text(None if exit_row is None else exit_row["status"]),
            len(per_chain),
            "是" if wrap_chains else "否",
            "、".join(wrap_chains),
            "、".join(native_only),
            mult.get("network"),
            mult.get("current"),
            mult.get("pending"),
            mult.get("error"),
            row["min_order_fiat_value"],
        ]
        for chain in chains:
            forms = per_chain.get(chain)
            if not forms:
                values.append("")
            elif forms == {"native", "wrapper_v2"}:
                values.append("原生+wrapper")
            elif "wrapper_v2" in forms:
                values.append("仅wrapper")
            else:
                values.append("仅原生")
        stock.append(values)
        status_cell = stock.cell(stock.max_row, 8)
        status_cell.fill = status_fill(None if exit_row is None else exit_row["status"])
        if wrap_chains:
            stock.cell(stock.max_row, 11).fill = WRAP_FILL
    style_header(stock, len(stock_headers))
    widths(stock, [16, 12, 10, 28, 16, 16, 16, 14, 36, 12, 12, 28, 28, 20, 22, 20, 24, 16]
           + [16] * len(chains))

    kraken = wb.create_sheet("Kraken现货")
    kraken_headers = [
        "交易对", "发行方符号", "正股代码", "上市地", "代币名称",
        "状态", "能不能直接卖给买盘", "最小下单量", "报价币", "充值网络",
        "当前multiplier", "multiplier查询网络",
    ]
    kraken.append(kraken_headers)
    for exit_row in sorted(exits, key=lambda row: (0 if row["status"] == "online" else 1, row["pair"])):
        symbol = exit_row["issuer_symbol"]
        underlying = next((row for row in under if row["issuer_symbol"] == symbol), None)
        mult = cache.get(symbol) or {}
        kraken.append([
            exit_row["pair"],
            symbol,
            None if underlying is None else underlying["underlying_symbol"],
            None if underlying is None else underlying["listing_country"],
            None if underlying is None else underlying["name"],
            exit_row["status"],
            sell_text(exit_row["status"]),
            exit_row["ordermin"],
            exit_row["quote"],
            exit_row["deposit_networks"],
            mult.get("current"),
            mult.get("network"),
        ])
        kraken.cell(kraken.max_row, 6).fill = status_fill(exit_row["status"])
    style_header(kraken, len(kraken_headers))
    widths(kraken, [16, 16, 12, 10, 28, 14, 36, 14, 12, 16, 22, 20])

    detail = wb.create_sheet("部署明细")
    detail_headers = [
        "发行方符号", "正股代码", "上市地", "链", "链ID", "形态", "合约地址",
        "支持原子兑换", "Kraken交易对", "Kraken状态",
    ]
    detail.append(detail_headers)
    under_by = {row["issuer_symbol"]: row for row in under}
    for dep in deps:
        underlying = under_by.get(dep["issuer_symbol"])
        exit_row = by_symbol_exit.get(dep["issuer_symbol"])
        atomic = dep["atomic_swaps"]
        detail.append([
            dep["issuer_symbol"],
            None if underlying is None else underlying["underlying_symbol"],
            None if underlying is None else underlying["listing_country"],
            dep["network"],
            dep["chain_id"],
            "wrapper" if dep["form"] == "wrapper_v2" else "原生",
            dep["address"],
            "" if atomic is None else ("是" if atomic else "否"),
            None if exit_row is None else exit_row["pair"],
            None if exit_row is None else exit_row["status"],
        ])
        if exit_row is not None:
            detail.cell(detail.max_row, 10).fill = status_fill(exit_row["status"])
    style_header(detail, len(detail_headers))
    widths(detail, [16, 12, 10, 18, 14, 12, 52, 16, 16, 14])

    wb.save(OUT)
    failed = sum(1 for row in cache.values() if row.get("error") and not row.get("current"))
    print(f"wrote {OUT} multipliers_failed={failed}", flush=True)
    return 0


if __name__ == "__main__":
    sys.exit(main())
